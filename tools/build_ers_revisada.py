from pathlib import Path
import json
import shutil

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\xampp\htdocs\Noble_Blend_Cafe")
DESKTOP = Path.home() / "Desktop"
OUT_DIR = ROOT / "artefatos_ers"
DIAGRAM_DIR = OUT_DIR / "diagramas"
SRC_DOCX = DESKTOP / "ERS cafeteria.docx"
OUT_DOCX = OUT_DIR / "ERS cafeteria - complemento 3.1 a 3.6.docx"


def font(size=24, bold=False):
    candidates = [
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def wrap_text(draw, text, fnt, max_width):
    words = text.split()
    lines = []
    line = ""
    for word in words:
        test = f"{line} {word}".strip()
        width = draw.textbbox((0, 0), test, font=fnt)[2]
        if width <= max_width:
            line = test
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def draw_box(draw, xy, title, body="", fill="#fff8ef", outline="#7b4a21"):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    title_font = font(26, True)
    body_font = font(20)
    draw.text((x1 + 18, y1 + 14), title, font=title_font, fill="#372416")
    if body:
        yy = y1 + 54
        for line in wrap_text(draw, body, body_font, x2 - x1 - 36):
            draw.text((x1 + 18, yy), line, font=body_font, fill="#372416")
            yy += 25


def draw_arrow(draw, start, end, label=""):
    draw.line([start, end], fill="#372416", width=3)
    x2, y2 = end
    draw.polygon([(x2, y2), (x2 - 12, y2 - 7), (x2 - 12, y2 + 7)], fill="#372416")
    if label:
        label_font = font(18)
        mx = (start[0] + end[0]) // 2
        my = (start[1] + end[1]) // 2 - 26
        draw.text((mx - 90, my), label, font=label_font, fill="#372416")


def make_sequence(path, title, participants, messages):
    width = 1600
    height = 260 + len(messages) * 70
    image = Image.new("RGB", (width, height), "#ffffff")
    draw = ImageDraw.Draw(image)
    title_font = font(36, True)
    part_font = font(22, True)
    msg_font = font(18)

    draw.text((60, 35), title, font=title_font, fill="#372416")
    margin = 110
    spacing = (width - 2 * margin) // (len(participants) - 1)
    xs = [margin + i * spacing for i in range(len(participants))]

    for x, name in zip(xs, participants):
        draw.rounded_rectangle((x - 95, 105, x + 95, 155), radius=12, fill="#f5e7d6", outline="#7b4a21", width=2)
        tw = draw.textbbox((0, 0), name, font=part_font)[2]
        draw.text((x - tw / 2, 118), name, font=part_font, fill="#372416")
        draw.line((x, 155, x, height - 40), fill="#bfa98f", width=2)

    y = 205
    idx = {name: i for i, name in enumerate(participants)}
    for src, dst, label in messages:
        x1 = xs[idx[src]]
        x2 = xs[idx[dst]]
        direction = 1 if x2 > x1 else -1
        draw.line((x1, y, x2, y), fill="#372416", width=3)
        draw.polygon([(x2, y), (x2 - 12 * direction, y - 7), (x2 - 12 * direction, y + 7)], fill="#372416")
        for i, line in enumerate(wrap_text(draw, label, msg_font, abs(x2 - x1) - 30)):
            draw.text((min(x1, x2) + 18, y - 28 + i * 22), line, font=msg_font, fill="#372416")
        y += 70

    image.save(path)


def make_conceptual(path):
    image = Image.new("RGB", (1900, 1250), "#ffffff")
    draw = ImageDraw.Draw(image)
    draw.text((60, 35), "Modelo Conceitual - Noble Blend Cafe", font=font(38, True), fill="#372416")

    boxes = {
        "Cliente": (90, 130, 450, 310, "nome, email, senha_hash, telefone, cpf, nascimento, possui_clube"),
        "Carrinho": (650, 130, 970, 285, "quantidade, preco_unitario"),
        "Produto": (1260, 130, 1660, 330, "nome, categoria, descricao, preco, preco_clube, estoque, imagem, ativo"),
        "Endereco": (90, 570, 470, 780, "destinatario, telefone, cep, endereco, numero, bairro, cidade, uf, frete"),
        "Pedido": (650, 570, 1030, 780, "numero_pedido, metodo_pagamento, subtotal, frete, desconto, total, status"),
        "ItemPedido": (1260, 570, 1660, 755, "nome_produto, quantidade, preco_unitario, subtotal"),
        "Funcionario": (650, 930, 1030, 1105, "nome, email, senha_hash, telefone, cargo"),
    }

    for title, (x1, y1, x2, y2, body) in boxes.items():
        draw_box(draw, (x1, y1, x2, y2), title, body)

    rel_font = font(20, True)
    relations = [
        ((450, 220), (650, 220), "mantem", "1", "N"),
        ((970, 220), (1260, 220), "referencia", "N", "1"),
        ((320, 310), (650, 620), "realiza", "1", "N"),
        ((470, 675), (650, 675), "entrega", "1", "N"),
        ((1030, 675), (1260, 675), "possui", "1", "N"),
        ((1460, 570), (1460, 330), "refere-se a", "N", "1"),
        ((840, 930), (840, 780), "atualiza status", "1", "N"),
    ]
    for start, end, label, start_card, end_card in relations:
        draw_arrow(draw, start, end)
        mx = (start[0] + end[0]) // 2
        my = (start[1] + end[1]) // 2
        tw = draw.textbbox((0, 0), label, font=rel_font)[2]
        draw.rectangle((mx - tw // 2 - 8, my - 34, mx + tw // 2 + 8, my - 6), fill="#ffffff")
        draw.text((mx - tw // 2, my - 34), label, font=rel_font, fill="#5b3a1f")
        draw.text((start[0] + 10, start[1] - 28), start_card, font=rel_font, fill="#5b3a1f")
        draw.text((end[0] - 28, end[1] + 8), end_card, font=rel_font, fill="#5b3a1f")

    image.save(path)


def add_after(paragraph, text=None, style=None, bold=False):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = paragraph._parent.add_paragraph()
    inserted._p = new_p
    if style:
        inserted.style = style
    if text:
        run = inserted.add_run(text)
        run.bold = bold
    return inserted


def insert_block_after_heading(doc, heading_text, lines):
    target = None
    for p in doc.paragraphs:
        if p.text.strip().lower() == heading_text.lower():
            target = p
            break
    if target is None:
        return

    current = target
    for item in reversed(lines):
        pass
    for item in lines:
        if isinstance(item, tuple) and item[0] == "image":
            current = add_after(current)
            current.alignment = WD_ALIGN_PARAGRAPH.CENTER
            current.add_run().add_picture(str(item[1]), width=Inches(item[2]))
        elif isinstance(item, tuple) and item[0] == "caption":
            current = add_after(current, item[1])
            current.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in current.runs:
                run.italic = True
                run.font.size = Pt(9)
        elif isinstance(item, tuple) and item[0] == "heading":
            current = add_after(current, item[1], style="Heading 3")
        else:
            current = add_after(current, str(item))
            for run in current.runs:
                run.font.size = Pt(10.5)


def make_report():
    report = OUT_DIR / "analise_coerencia_ers.md"
    text = """# Analise de coerencia da ERS - Noble Blend Cafe

## Funcionalidades implementadas no sistema atual
- Cadastro, login e perfil de cliente.
- Cadastro, login e perfil de funcionario.
- Cardapio com busca, categoria, preco maximo, ordenacao e filtro de estoque.
- Carrinho com adicionar, atualizar quantidade e remover item.
- Checkout com endereco de entrega, consulta de CEP via ViaCEP, pagamento simulado por PIX/credito/debito, frete e desconto PIX.
- Confirmacao e historico de pedidos do cliente.
- Painel do funcionario com metricas, produtos mais vendidos, receita por categoria, pedidos por status e pagamentos.
- Gestao de produtos/cardapio, incluindo preco normal, preco Coffee Lovers, estoque, imagem e status ativo.
- Listagem de clientes, funcionarios e pedidos.
- Atualizacao de status do pedido pelo funcionario.
- Coffee Lovers simples: ativacao por CPF/senha e uso do preco de clube.

## Pontos do documento que precisam ajuste
- Personalizacao por tamanho, tipo de leite e adicionais ainda nao existe na tela de cardapio/carrinho.
- Agendamento de pedidos ainda nao existe no checkout.
- Retirada no local nao aparece como opcao; o fluxo atual e entrega em domicilio.
- Area de entrega nao e validada; o sistema consulta CEP via ViaCEP e grava endereco.
- Pagamento nao tem TEF/gateway real; esta simulado.
- Pontos, itens gratuitos e beneficios acumulados do Coffee Lovers ainda nao existem; existe desconto por preco de clube.
- Cancelamento de pedido pelo cliente ainda nao existe; cancelamento aparece como status editavel pelo funcionario.
- Relatorios existem parcialmente no painel do funcionario, mas nao ha filtro por periodo nem relatorio formal exportavel.

## Telas indicadas para prints
- Login do cliente: view/login_cliente.php
- Cadastro de cliente: view/cadastro_cliente.php
- Home do cliente: view/logado_cliente.php
- Cardapio e filtros: view/cardapio.php
- Carrinho: view/carrinho.php
- Checkout: view/checkout.php
- Confirmacao do pedido: view/confirmacao.php?pedido=NUMERO_DO_PEDIDO
- Meus pedidos: view/cliente_pedidos.php
- Coffee Lovers: view/cadastro_coffee_lovers.php
- Login do funcionario: view/login_funcionario.php
- Painel geral do funcionario: view/logado_funcionario.php
- Produtos/cardapio administrativo: view/funcionario_produtos.php
- Pedidos/status: view/funcionario_pedidos.php
- Clientes: view/funcionario_clientes.php
- Funcionarios: view/funcionario_funcionarios.php
"""
    report.write_text(text, encoding="utf-8")


def make_staruml_notes():
    puml = OUT_DIR / "diagramas_para_staruml.puml"
    puml.write_text(
        """@startuml
title DSS - Realizar Pedido Online
actor Cliente
participant Sistema
Cliente -> Sistema: autentica-se
Cliente -> Sistema: consulta cardapio e filtros
Cliente -> Sistema: adiciona produto ao carrinho
Cliente -> Sistema: revisa carrinho
Cliente -> Sistema: informa entrega e pagamento
Sistema --> Cliente: confirma pedido e numero
@enduml

@startuml
title DSS - Atender/Monitorar Pedido
actor Funcionario
participant Sistema
Funcionario -> Sistema: autentica-se
Funcionario -> Sistema: consulta painel de pedidos
Funcionario -> Sistema: altera status do pedido
Sistema --> Funcionario: registra novo status
@enduml

@startuml
title Modelo Conceitual
class Cliente
class Carrinho
class Produto
class Pedido
class Endereco
class ItemPedido
class Funcionario
Cliente "1" -- "N" Carrinho
Produto "1" -- "N" Carrinho
Cliente "1" -- "N" Pedido
Endereco "1" -- "N" Pedido
Pedido "1" -- "N" ItemPedido
Produto "1" -- "N" ItemPedido
Funcionario ..> Pedido : atualiza status
@enduml
""",
        encoding="utf-8",
    )

    mdj = OUT_DIR / "noble_blend_modelo_staruml_resumo.mdj"
    mdj.write_text(
        json.dumps(
            {
                "_type": "Project",
                "name": "Noble Blend Cafe - ERS ate 3.6",
                "ownedElements": [
                    {
                        "_type": "UMLModel",
                        "name": "Analise",
                        "documentation": "Projeto de apoio criado a partir do codigo e do banco SQL. Use o arquivo diagramas_para_staruml.puml como roteiro/importacao para os diagramas no StarUML.",
                    }
                ],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def build_docx():
    if SRC_DOCX.exists():
        shutil.copy2(SRC_DOCX, OUT_DOCX)
        doc = Document(OUT_DOCX)
    else:
        doc = Document()
        doc.add_heading("Noble Blend Cafe - Complemento da ERS ate o item 3.6", 0)
        doc.add_paragraph(
            "Documento complementar elaborado a partir do codigo-fonte e do banco SQL do sistema atual."
        )
        for heading in [
            "2.4  REQUISITOS ADIADOS",
            "3. REQUISITOS ESPECÍFICOS",
            "3.1.1 Interfaces do Usuário dos Casos de Uso (Inicial maiúsculo e negrito)",
            "3.1.2 Interfaces do Sistema",
            "3.1.3 Interfaces de Hardware",
            "3.1.4 Interfaces de Software",
            "3.1.5 Interfaces de Comunicação",
            "3.5  DIAGRAMAS DE SEQUÊNCIA DE EVENTOS DO SISTEMA",
            "3.6  MODELO CONCEITUAL",
        ]:
            doc.add_heading(heading, 1 if heading.startswith(("2.", "3.")) and "  " not in heading[4:] else 2)

    styles = doc.styles
    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        if style_name in styles:
            styles[style_name].font.name = "Arial"

    intro_lines = [
        "Analise de coerencia com o sistema atual: o projeto implementado cobre cadastro/autenticacao de clientes e funcionarios, cardapio, carrinho, checkout com entrega, pedidos, atualizacao de status, gestao de produtos, controle de estoque e adesao simples ao Coffee Lovers. Recursos como personalizacao por tamanho/tipo de leite/adicionais, agendamento, retirada no local, pontos acumulados e gateway/TEF real devem ser tratados como requisitos adiados ou futuros.",
    ]
    insert_block_after_heading(doc, "3. REQUISITOS ESPECÍFICOS", intro_lines)

    insert_block_after_heading(
        doc,
        "3.1.1 Interfaces do Usuário dos Casos de Uso (Inicial maiúsculo e negrito)",
        [
            "As interfaces abaixo correspondem diretamente as telas existentes no sistema Web Noble Blend Cafe e podem ser usadas como evidencias visuais dos casos de uso.",
            "Realizar pedido online: login_cliente.php, logado_cliente.php, cardapio.php, carrinho.php, checkout.php e confirmacao.php. Prints recomendados: Cardapio com filtros, Carrinho com resumo e Checkout com entrega/pagamento.",
            "Aderir ao Clube Coffee Lovers: cadastro_coffee_lovers.php. Print recomendado: formulario de CPF/senha e, depois da ativacao, tela de cadastro concluido.",
            "Monitorar/atender pedido: login_funcionario.php, logado_funcionario.php e funcionario_pedidos.php. Prints recomendados: Painel geral com metricas e tela Pedidos com seletor de status.",
            "Gerenciar produtos e estoque: funcionario_produtos.php. Print recomendado: formulario de cadastro/edicao de produto e tabela com preco, preco clube e estoque.",
            "Consultar clientes/funcionarios: funcionario_clientes.php e funcionario_funcionarios.php. Prints recomendados: listagens administrativas.",
            "Manter perfil: perfil_cliente.php e funcionario_perfil.php. Print recomendado: formulario de edicao de dados cadastrais.",
        ],
    )

    insert_block_after_heading(
        doc,
        "3.1.2 Interfaces do Sistema",
        [
            "O sistema interage com o banco de dados MySQL/MariaDB noble_blend_cafe por meio da classe Conexao.php. As principais tabelas utilizadas sao clientes, funcionarios, produtos, carrinho, enderecos, pedidos e pedido_itens.",
            "Na tela de checkout, o sistema consulta o servico ViaCEP por requisicao HTTP para buscar rua, bairro, cidade e UF a partir do CEP informado pelo cliente.",
            "O pagamento esta implementado como simulacao interna de metodo de pagamento, permitindo PIX, cartao de credito e cartao de debito. Nao ha integracao real com TEF, adquirente ou gateway de pagamento nesta versao.",
        ],
    )

    insert_block_after_heading(
        doc,
        "3.1.3 Interfaces de Hardware",
        [
            "O produto nao depende de hardware especifico. A utilizacao ocorre por computadores, notebooks, tablets ou smartphones com navegador Web.",
            "Para o uso administrativo, funcionarios precisam apenas de dispositivo com acesso ao navegador e conexao ao servidor do sistema. Leitores biometricos, impressoras fiscais e maquininhas TEF nao estao integrados nesta versao.",
        ],
    )

    insert_block_after_heading(
        doc,
        "3.1.4 Interfaces de Software",
        [
            "O sistema e uma aplicacao Web desenvolvida em PHP, HTML, CSS e JavaScript, executada em ambiente Apache/PHP, como o XAMPP usado no desenvolvimento.",
            "O banco de dados utilizado e MySQL/MariaDB, com script de criacao em Banco/noble_blend_cafe.sql.",
            "O navegador do usuario deve suportar HTML5, CSS3, JavaScript e requisicoes fetch para consulta de CEP.",
            "O servidor PHP utiliza sessoes para manter o estado de autenticacao de clientes e funcionarios.",
        ],
    )

    insert_block_after_heading(
        doc,
        "3.1.5 Interfaces de Comunicação",
        [
            "A comunicacao entre navegador e servidor ocorre por HTTP/HTTPS, com formularios POST/GET para cadastro, login, carrinho, checkout, produtos e pedidos.",
            "A comunicacao entre PHP e banco de dados ocorre por conexao MySQL usando mysqli e consultas preparadas.",
            "A consulta externa de CEP ocorre via HTTPS para o endpoint publico do ViaCEP, retornando dados em JSON.",
        ],
    )

    pedido_png = DIAGRAM_DIR / "dss_realizar_pedido.png"
    coffee_png = DIAGRAM_DIR / "dss_coffee_lovers.png"
    monitor_png = DIAGRAM_DIR / "dss_monitorar_pedidos.png"
    produto_png = DIAGRAM_DIR / "dss_gerenciar_produtos.png"
    conceitual_png = DIAGRAM_DIR / "modelo_conceitual.png"

    insert_block_after_heading(
        doc,
        "3.5  DIAGRAMAS DE SEQUÊNCIA DE EVENTOS DO SISTEMA",
        [
            ("image", pedido_png, 6.4),
            ("caption", "Figura - Diagrama de sequencia do sistema: Realizar Pedido Online."),
            ("image", coffee_png, 6.4),
            ("caption", "Figura - Diagrama de sequencia do sistema: Aderir ao Clube Coffee Lovers."),
            ("image", monitor_png, 6.4),
            ("caption", "Figura - Diagrama de sequencia do sistema: Monitorar/Atender Pedidos."),
            ("image", produto_png, 6.4),
            ("caption", "Figura - Diagrama de sequencia do sistema: Gerenciar Produtos e Estoque."),
        ],
    )

    insert_block_after_heading(
        doc,
        "3.6  MODELO CONCEITUAL",
        [
            "O modelo conceitual foi elaborado a partir das entidades persistidas no banco de dados do sistema atual. O Cliente pode manter itens no Carrinho e realizar Pedidos. Cada Pedido possui um Endereco de entrega e diversos Itens de Pedido. Os itens referenciam Produtos, e o Funcionario atua na manutencao de produtos e atualizacao do status dos pedidos.",
            ("image", conceitual_png, 6.5),
            ("caption", "Figura - Modelo conceitual do sistema Noble Blend Cafe."),
        ],
    )

    insert_block_after_heading(
        doc,
        "2.4  REQUISITOS ADIADOS",
        [
            "Revisao de coerencia: alem de HU_S03 e HU_S04, tambem devem ser considerados adiados ou futuros os recursos de agendamento de pedido, retirada no local, personalizacao por tamanho/tipo de leite/adicionais, beneficios por pontos/itens gratuitos e integracao real de pagamento/TEF. A versao atual possui pagamento simulado, entrega com endereco, desconto PIX e preco Coffee Lovers.",
        ],
    )

    doc.save(OUT_DOCX)


def main():
    OUT_DIR.mkdir(exist_ok=True)
    DIAGRAM_DIR.mkdir(exist_ok=True)

    make_sequence(
        DIAGRAM_DIR / "dss_realizar_pedido.png",
        "DSS - Realizar Pedido Online",
        ["Cliente", "Sistema", "Banco"],
        [
            ("Cliente", "Sistema", "Autentica-se e acessa o cardapio"),
            ("Cliente", "Sistema", "Filtra produtos e escolhe quantidade"),
            ("Sistema", "Banco", "Consulta produtos ativos e estoque"),
            ("Cliente", "Sistema", "Adiciona item ao carrinho"),
            ("Sistema", "Banco", "Grava/atualiza item do carrinho"),
            ("Cliente", "Sistema", "Informa entrega e metodo de pagamento"),
            ("Sistema", "Banco", "Cria endereco, pedido e itens; baixa estoque"),
            ("Sistema", "Cliente", "Exibe numero e confirmacao do pedido"),
        ],
    )
    make_sequence(
        DIAGRAM_DIR / "dss_coffee_lovers.png",
        "DSS - Aderir ao Clube Coffee Lovers",
        ["Cliente", "Sistema", "Banco"],
        [
            ("Cliente", "Sistema", "Acessa cadastro Coffee Lovers"),
            ("Cliente", "Sistema", "Informa CPF, senha e aceite de marketing"),
            ("Sistema", "Banco", "Valida senha do cliente"),
            ("Sistema", "Banco", "Ativa possui_clube e coffee_lover"),
            ("Sistema", "Banco", "Atualiza precos do carrinho para preco_clube"),
            ("Sistema", "Cliente", "Confirma cadastro concluido"),
        ],
    )
    make_sequence(
        DIAGRAM_DIR / "dss_monitorar_pedidos.png",
        "DSS - Monitorar/Atender Pedidos",
        ["Funcionario", "Sistema", "Banco"],
        [
            ("Funcionario", "Sistema", "Autentica-se na area administrativa"),
            ("Sistema", "Banco", "Consulta pedidos recentes e metricas"),
            ("Sistema", "Funcionario", "Exibe painel e lista de pedidos"),
            ("Funcionario", "Sistema", "Seleciona novo status do pedido"),
            ("Sistema", "Banco", "Atualiza status se permitido"),
            ("Sistema", "Funcionario", "Exibe confirmacao da atualizacao"),
        ],
    )
    make_sequence(
        DIAGRAM_DIR / "dss_gerenciar_produtos.png",
        "DSS - Gerenciar Produtos e Estoque",
        ["Funcionario", "Sistema", "Banco"],
        [
            ("Funcionario", "Sistema", "Acessa tela Produtos"),
            ("Sistema", "Banco", "Lista produtos ativos e inativos"),
            ("Funcionario", "Sistema", "Cadastra ou edita produto"),
            ("Sistema", "Banco", "Valida dados e salva produto/estoque"),
            ("Funcionario", "Sistema", "Remove produto do cardapio"),
            ("Sistema", "Banco", "Marca produto como inativo"),
        ],
    )
    make_conceptual(DIAGRAM_DIR / "modelo_conceitual.png")
    make_report()
    make_staruml_notes()
    build_docx()
    print(str(OUT_DOCX))
    print(str(OUT_DIR))


if __name__ == "__main__":
    main()
